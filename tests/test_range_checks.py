"""Checks 33 and 34 — the two findings T6 showed were reachable structurally.

Both catch errors that are invisible to every other check and to an evaluator
alike: the arithmetic is performed exactly as written, and the answer is wrong
anyway.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from tickmark.checks.double_counting import DoubleCountingCheck
from tickmark.checks.short_range import ShortRangeCheck
from tickmark.findings.model import Severity
from tickmark.workbook import open_workbook


def build(tmp_path: Path, cells: dict[str, object], name: str = "r.xlsx") -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for coord, value in cells.items():
        ws[coord] = value
    path = tmp_path / name
    wb.save(path)
    return path


def run(check, path: Path):
    with open_workbook(path) as wb:
        return check.run(wb.sheets[0])


# ------------------------------------------------------------ check 33


class TestShortRange:
    def test_total_misses_rows_added_below(self, tmp_path: Path):
        # The classic: data grew to B13, the SUM was never extended past B12.
        cells = {f"B{r}": r * 100 for r in range(2, 14)}
        cells["B15"] = "=SUM(B2:B12)"
        found = run(ShortRangeCheck(), build(tmp_path, cells))

        assert len(found) == 1
        assert found[0].severity is Severity.HIGH
        assert found[0].context["missed"] == "B13"
        assert "continues at B13" in found[0].explanation

    def test_several_missed_rows_are_reported_as_a_span(self, tmp_path: Path):
        cells = {f"B{r}": r * 100 for r in range(2, 20)}
        cells["B21"] = "=SUM(B2:B12)"
        found = run(ShortRangeCheck(), build(tmp_path, cells))

        assert found[0].context["missed"] == "B13:B19"
        assert found[0].context["missed_count"] == "7"

    def test_upward_misses_are_deliberately_not_reported(self, tmp_path: Path):
        # A range can also start too low, but above a block sit headers, section
        # titles and numeric labels that no local test distinguishes from data.
        # Reporting the ambiguous direction would cost this check its credibility.
        # Deferred to v1.1 by decision, pinned here so it is not "fixed" silently.
        cells = {"A1": "Header"}
        cells.update({f"B{r}": r * 100 for r in range(2, 14)})
        cells["B15"] = "=SUM(B5:B13)"  # starts too low
        assert run(ShortRangeCheck(), build(tmp_path, cells)) == []

    def test_aggregates_other_than_sum(self, tmp_path: Path):
        for function in ("AVERAGE", "MAX", "COUNT", "MEDIAN"):
            cells = {f"B{r}": r * 100 for r in range(2, 14)}
            cells["B15"] = f"={function}(B2:B12)"
            found = run(ShortRangeCheck(), build(tmp_path, cells, f"{function}.xlsx"))
            assert len(found) == 1, function


class TestShortRangeSuppressions:
    def test_a_correct_total_is_silent(self, tmp_path: Path):
        cells = {f"B{r}": r * 100 for r in range(2, 13)}
        cells["B13"] = "=SUM(B2:B12)"
        assert run(ShortRangeCheck(), build(tmp_path, cells)) == []

    def test_the_formula_cell_itself_is_not_a_missing_row(self, tmp_path: Path):
        # =SUM(B2:B12) in B13 must not be told to include B13 — that is a
        # circular reference, and check 5's job.
        cells = {f"B{r}": r * 100 for r in range(2, 13)}
        cells["B13"] = "=SUM(B2:B12)"
        found = run(ShortRangeCheck(), build(tmp_path, cells))
        assert all("B13" not in f.context.get("missed", "") for f in found)

    def test_a_grand_total_below_a_subtotal_is_not_a_missing_row(self, tmp_path: Path):
        # Telling the user to include the total row would be advice to
        # double-count — the exact bug check 34 reports.
        cells = {f"B{r}": r * 100 for r in range(2, 11)}
        cells["B11"] = "=SUM(B2:B10)"
        cells["D1"] = "=SUM(B2:B10)"
        found = run(ShortRangeCheck(), build(tmp_path, cells))
        assert found == []

    def test_text_ends_the_block(self, tmp_path: Path):
        cells = {f"B{r}": r * 100 for r in range(2, 13)}
        cells["B13"] = "Total"
        cells["B15"] = "=SUM(B2:B12)"
        assert run(ShortRangeCheck(), build(tmp_path, cells)) == []

    def test_a_blank_ends_the_block(self, tmp_path: Path):
        cells = {f"B{r}": r * 100 for r in range(2, 13)}
        cells["B14"] = 999  # a gap at B13 separates this deliberately
        cells["B16"] = "=SUM(B2:B12)"
        assert run(ShortRangeCheck(), build(tmp_path, cells)) == []

    def test_a_column_of_data_under_a_text_header_still_reports(self, tmp_path: Path):
        # Regression: header_rows once classified every formula-free row as a
        # header, so a plain column of numbers with one SUM at the bottom made
        # the whole block "header" and this check went silent on its own use case.
        cells = {"B1": "Amount"}
        cells.update({f"B{r}": r * 100 for r in range(2, 14)})
        cells["B15"] = "=SUM(B2:B12)"
        found = run(ShortRangeCheck(), build(tmp_path, cells))
        assert [f.context["missed"] for f in found] == ["B13"]

    def test_whole_column_range_cannot_stop_short(self, tmp_path: Path):
        cells = {f"B{r}": r * 100 for r in range(2, 14)}
        cells["D1"] = "=SUM(B:B)"
        assert run(ShortRangeCheck(), build(tmp_path, cells)) == []

    def test_rectangular_ranges_are_left_alone(self, tmp_path: Path):
        # Four edges and many legitimate reasons to stop; guessing costs more in
        # false positives than it earns.
        cells = {}
        for col in "BCD":
            for row in range(2, 14):
                cells[f"{col}{row}"] = row
        cells["F1"] = "=SUM(B2:C12)"
        assert run(ShortRangeCheck(), build(tmp_path, cells)) == []

    def test_cross_sheet_ranges_are_skipped(self, tmp_path: Path):
        wb = Workbook()
        data = wb.active
        data.title = "Data"
        for row in range(2, 14):
            data[f"B{row}"] = row
        summary = wb.create_sheet("Summary")
        summary["A1"] = "=SUM(Data!B2:B12)"
        path = tmp_path / "cross.xlsx"
        wb.save(path)

        with open_workbook(path) as book:
            assert ShortRangeCheck().run(book.sheet("Summary")) == []


# ------------------------------------------------------------ check 34


class TestDoubleCounting:
    def test_subtotal_inside_a_grand_total(self, tmp_path: Path):
        cells = {f"B{r}": r * 100 for r in range(2, 10)}
        cells["B10"] = "=SUM(B2:B9)"  # subtotal
        cells.update({f"B{r}": r * 100 for r in range(11, 20)})
        cells["B21"] = "=SUM(B2:B20)"  # counts B2:B9 twice

        found = run(DoubleCountingCheck(), build(tmp_path, cells))
        assert len(found) == 1
        assert found[0].severity is Severity.HIGH
        assert found[0].context["subtotal_cell"] == "B10"
        assert "counted twice" in found[0].explanation

    def test_reports_which_range_is_duplicated(self, tmp_path: Path):
        cells = {f"B{r}": r for r in range(2, 10)}
        cells["B10"] = "=SUM(B2:B9)"
        cells["B21"] = "=SUM(B2:B20)"
        found = run(DoubleCountingCheck(), build(tmp_path, cells))
        assert found[0].context["subtotal_range"] == "B2:B9"
        assert found[0].context["range"] == "B2:B20"

    def test_average_is_also_distorted(self, tmp_path: Path):
        cells = {f"B{r}": r for r in range(2, 10)}
        cells["B10"] = "=SUM(B2:B9)"
        cells["B21"] = "=AVERAGE(B2:B20)"
        assert len(run(DoubleCountingCheck(), build(tmp_path, cells))) == 1


class TestDoubleCountingSuppressions:
    def test_a_cell_sourced_from_elsewhere_is_normal(self, tmp_path: Path):
        # B10 draws from column D, so including it double-counts nothing. This
        # is the common, correct pattern the check must not flag.
        cells = {f"B{r}": r for r in range(2, 10)}
        cells["B10"] = "=SUM(D1:D5)"
        cells.update({f"D{r}": r for r in range(1, 6)})
        cells["B21"] = "=SUM(B2:B20)"
        assert run(DoubleCountingCheck(), build(tmp_path, cells)) == []

    def test_subtotal_outside_the_range_is_fine(self, tmp_path: Path):
        cells = {f"B{r}": r for r in range(2, 10)}
        cells["B30"] = "=SUM(B2:B9)"  # outside the grand total's range
        cells["B21"] = "=SUM(B2:B20)"
        assert run(DoubleCountingCheck(), build(tmp_path, cells)) == []

    def test_self_inclusion_is_left_to_check_five(self, tmp_path: Path):
        # =SUM(B2:B20) living inside its own range is a circular reference.
        cells = {f"B{r}": r for r in range(2, 10)}
        cells["B10"] = "=SUM(B2:B20)"
        assert run(DoubleCountingCheck(), build(tmp_path, cells)) == []

    def test_a_plain_total_column_is_silent(self, tmp_path: Path):
        cells = {f"B{r}": r for r in range(2, 20)}
        cells["B21"] = "=SUM(B2:B20)"
        assert run(DoubleCountingCheck(), build(tmp_path, cells)) == []

    def test_nested_non_overlapping_subtotals_are_fine(self, tmp_path: Path):
        # Two subtotals over disjoint blocks, added by name rather than by range.
        cells = {f"B{r}": r for r in range(2, 10)}
        cells.update({f"B{r}": r for r in range(11, 19)})
        cells["B10"] = "=SUM(B2:B9)"
        cells["B19"] = "=SUM(B11:B18)"
        cells["B21"] = "=B10+B19"
        assert run(DoubleCountingCheck(), build(tmp_path, cells)) == []


@pytest.mark.parametrize("check", [ShortRangeCheck(), DoubleCountingCheck()])
def test_unparseable_formula_does_not_crash(tmp_path: Path, check):
    cells = {f"B{r}": r for r in range(2, 13)}
    cells["B14"] = '=SUM(B2:B12,"unterminated'
    assert isinstance(run(check, build(tmp_path, cells)), list)
