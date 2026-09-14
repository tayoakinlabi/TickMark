"""The dependency graph, and checks 5 and 7 built on it."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from tickmark.checks.circular_refs import CircularRefsCheck
from tickmark.checks.complexity import ComplexityCheck
from tickmark.findings.model import Severity
from tickmark.graph.complexity import score_formula
from tickmark.graph.cycles import find_cycles
from tickmark.graph.dependency import CellKey, DependencyGraph
from tickmark.workbook import open_workbook


def build(tmp_path: Path, sheets: dict[str, dict[str, object]], name: str = "g.xlsx") -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    for sheet_name, cells in sheets.items():
        ws = wb.create_sheet(sheet_name)
        for coord, value in cells.items():
            ws[coord] = value
    path = tmp_path / name
    wb.save(path)
    return path


def graph_for(path: Path) -> DependencyGraph:
    with open_workbook(path) as wb:
        return DependencyGraph(wb)


class TestDependencyGraph:
    def test_precedents_link_formula_cells(self, tmp_path: Path):
        path = build(tmp_path, {"S": {"A1": 1, "B1": "=A1*2", "C1": "=B1+1"}})
        graph = graph_for(path)
        assert CellKey("S", 1, 2) in graph.precedents(CellKey("S", 1, 3))

    def test_constants_are_not_nodes(self, tmp_path: Path):
        # A1 holds a literal, so it cannot be part of a loop and is not a node.
        path = build(tmp_path, {"S": {"A1": 1, "B1": "=A1*2"}})
        graph = graph_for(path)
        assert CellKey("S", 1, 1) not in graph.cells

    def test_dependents_are_the_inverse(self, tmp_path: Path):
        path = build(tmp_path, {"S": {"A1": 1, "B1": "=A1*2", "C1": "=B1+1"}})
        graph = graph_for(path)
        assert graph.dependents(CellKey("S", 1, 2)) == {CellKey("S", 1, 3)}

    def test_ranges_link_every_formula_inside(self, tmp_path: Path):
        cells = {f"A{r}": f"=B{r}*2" for r in range(1, 6)}
        cells["C1"] = "=SUM(A1:A5)"
        path = build(tmp_path, {"S": cells})
        graph = graph_for(path)
        assert len(graph.precedents(CellKey("S", 1, 3))) == 5

    def test_cross_sheet_references(self, tmp_path: Path):
        path = build(
            tmp_path,
            {"Data": {"A1": 1, "B1": "=A1*2"}, "Summary": {"A1": "=Data!B1+1"}},
        )
        graph = graph_for(path)
        assert CellKey("Data", 1, 2) in graph.precedents(CellKey("Summary", 1, 1))

    def test_whole_column_reference_is_not_expanded(self, tmp_path: Path):
        # =SUM(A:A) names a million cells. Only the formulas in it become edges,
        # and the run must complete quickly.
        cells = {f"A{r}": f"=B{r}*2" for r in range(1, 21)}
        cells["C1"] = "=SUM(A:A)"
        path = build(tmp_path, {"S": cells})
        graph = graph_for(path)
        assert len(graph.precedents(CellKey("S", 1, 3))) == 20

    def test_external_references_create_no_edges(self, tmp_path: Path):
        path = build(tmp_path, {"S": {"A1": "=[1]Other!B2+1"}})
        graph = graph_for(path)
        assert graph.precedents(CellKey("S", 1, 1)) == frozenset()


class TestCycles:
    def test_self_referencing_sum_is_a_cycle(self, tmp_path: Path):
        # The classic: a total typed one row too far down.
        path = build(tmp_path, {"S": {"A1": 1, "A2": 2, "A3": "=SUM(A1:A3)"}})
        cycles = find_cycles(graph_for(path))
        assert cycles == [[CellKey("S", 3, 1)]]

    def test_two_cell_loop(self, tmp_path: Path):
        path = build(tmp_path, {"S": {"A1": "=B1+1", "B1": "=A1+1"}})
        cycles = find_cycles(graph_for(path))
        assert len(cycles) == 1
        assert len(cycles[0]) == 2

    def test_three_cell_loop(self, tmp_path: Path):
        path = build(tmp_path, {"S": {"A1": "=C1+1", "B1": "=A1+1", "C1": "=B1+1"}})
        cycles = find_cycles(graph_for(path))
        assert len(cycles) == 1
        assert len(cycles[0]) == 3

    def test_loop_across_sheets(self, tmp_path: Path):
        path = build(tmp_path, {"One": {"A1": "=Two!A1+1"}, "Two": {"A1": "=One!A1+1"}})
        cycles = find_cycles(graph_for(path))
        assert len(cycles) == 1
        assert {k.sheet for k in cycles[0]} == {"One", "Two"}

    def test_a_long_chain_is_not_a_cycle(self, tmp_path: Path):
        cells = {"A1": 1}
        cells.update({f"A{r}": f"=A{r - 1}+1" for r in range(2, 60)})
        assert find_cycles(graph_for(build(tmp_path, {"S": cells}))) == []

    def test_a_deep_chain_does_not_hit_the_recursion_limit(self, tmp_path: Path):
        # Written iteratively for exactly this reason.
        cells = {"A1": 1}
        cells.update({f"A{r}": f"=A{r - 1}+1" for r in range(2, 3000)})
        assert find_cycles(graph_for(build(tmp_path, {"S": cells}))) == []

    def test_clean_workbook_has_no_cycles(self, tmp_path: Path):
        path = build(tmp_path, {"S": {"A1": 1, "B1": "=A1*2", "C1": "=B1+1"}})
        assert find_cycles(graph_for(path)) == []


class TestCircularRefsCheck:
    def test_self_reference_is_explained_plainly(self, tmp_path: Path):
        path = build(tmp_path, {"S": {"A1": 1, "A2": 2, "A3": "=SUM(A1:A3)"}})
        with open_workbook(path) as wb:
            found = CircularRefsCheck().run_workbook(wb)
        assert len(found) == 1
        assert found[0].severity is Severity.HIGH
        assert "typed one row short" in found[0].explanation

    def test_multi_cell_loop_names_the_path(self, tmp_path: Path):
        path = build(tmp_path, {"S": {"A1": "=C1+1", "B1": "=A1+1", "C1": "=B1+1"}})
        with open_workbook(path) as wb:
            found = CircularRefsCheck().run_workbook(wb)
        assert "→" in found[0].context["path"]
        assert found[0].context["cycle_length"] == "3"

    def test_clean_workbook_is_silent(self, tmp_path: Path):
        path = build(tmp_path, {"S": {"A1": 1, "B1": "=A1*2"}})
        with open_workbook(path) as wb:
            assert CircularRefsCheck().run_workbook(wb) == []


class TestComplexity:
    def test_simple_formula_is_not_notable(self, tmp_path: Path):
        path = build(tmp_path, {"S": {"A1": 1, "B1": "=A1*2"}})
        with open_workbook(path) as wb:
            cell = next(wb.sheets[0].formulas())
            assert score_formula(cell).is_notable is False

    def test_deep_nesting_is_notable(self, tmp_path: Path):
        deep = "=IF(A1>1,IF(A1>2,IF(A1>3,IF(A1>4,IF(A1>5,IF(A1>6,7,6),5),4),3),2),1)"
        path = build(tmp_path, {"S": {"B1": deep}})
        with open_workbook(path) as wb:
            score = score_formula(next(wb.sheets[0].formulas()))
        assert score.is_notable
        assert "nested" in score.reason

    def test_ranking_is_bounded(self, tmp_path: Path):
        deep = "=IF(A1>1,IF(A1>2,IF(A1>3,IF(A1>4,IF(A1>5,IF(A1>6,7,6),5),4),3),2),1)"
        cells = {f"B{r}": deep for r in range(1, 40)}
        path = build(tmp_path, {"S": cells})
        with open_workbook(path) as wb:
            found = ComplexityCheck(limit=5).run_workbook(wb)
        assert len(found) == 5

    def test_says_how_many_were_omitted(self, tmp_path: Path):
        deep = "=IF(A1>1,IF(A1>2,IF(A1>3,IF(A1>4,IF(A1>5,IF(A1>6,7,6),5),4),3),2),1)"
        cells = {f"B{r}": deep for r in range(1, 40)}
        path = build(tmp_path, {"S": cells})
        with open_workbook(path) as wb:
            found = ComplexityCheck(limit=5).run_workbook(wb)
        assert "39 formulas" in found[0].explanation

    def test_tidy_workbook_reports_nothing(self, tmp_path: Path):
        cells = {f"B{r}": f"=A{r}*2" for r in range(1, 30)}
        path = build(tmp_path, {"S": cells})
        with open_workbook(path) as wb:
            assert ComplexityCheck().run_workbook(wb) == []

    @pytest.mark.parametrize("severity_is_info", [True])
    def test_complexity_is_informational_not_an_error(self, tmp_path: Path, severity_is_info: bool):
        # Nothing here is wrong; it is a maintenance note.
        deep = "=IF(A1>1,IF(A1>2,IF(A1>3,IF(A1>4,IF(A1>5,IF(A1>6,7,6),5),4),3),2),1)"
        path = build(tmp_path, {"S": {"B1": deep}})
        with open_workbook(path) as wb:
            found = ComplexityCheck().run_workbook(wb)
        assert found[0].severity is Severity.INFO
