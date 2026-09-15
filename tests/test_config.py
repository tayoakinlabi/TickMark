"""``tickmark.toml`` — T4's answer, and item 19's user-defined rules.

The theme running through these tests is that a config file is edited by hand,
often in a hurry, by someone who is not a programmer. So the behaviour that
matters most is not that a correct file works — it is that an incorrect one
costs the user only the part they got wrong, says which part, and still produces
their audit.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from tickmark.checks.registry import run_audit
from tickmark.config.file import ConfigError, find_config, load_config
from tickmark.config.rules import DEFAULT_RULES
from tickmark.findings.model import Severity
from tickmark.workbook.loader import open_workbook


def write(path: Path, text: str) -> Path:
    path.write_text(text.strip() + "\n", encoding="utf-8")
    return path


@pytest.fixture
def workbook(tmp_path: Path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws["A1"] = "Item"
    ws["B1"] = "Net"
    for row in range(2, 8):
        ws[f"A{row}"] = f"item-{row}"
        ws[f"B{row}"] = row * 10
        ws[f"C{row}"] = f"=B{row}*1.075"
    ws["D2"] = "TODO: check this with finance"
    path = tmp_path / "book.xlsx"
    wb.save(path)
    return path


class TestDiscovery:
    def test_found_beside_the_workbook(self, tmp_path: Path):
        config = write(tmp_path / "tickmark.toml", "[rules]\ncomplexity_limit = 3")
        assert find_config(tmp_path / "book.xlsx") == config

    def test_found_by_walking_upwards(self, tmp_path: Path):
        # The shape this exists for: a shared drive of nested folders with one
        # config at the top.
        config = write(tmp_path / "tickmark.toml", "[rules]\ncomplexity_limit = 3")
        nested = tmp_path / "2024" / "Q3"
        nested.mkdir(parents=True)
        assert find_config(nested / "book.xlsx") == config

    def test_absent_is_not_an_error(self, tmp_path: Path):
        assert find_config(tmp_path / "book.xlsx") is None


class TestSettings:
    def test_settings_are_applied(self, tmp_path: Path):
        config = write(
            tmp_path / "tickmark.toml",
            """
            [rules]
            complexity_limit = 3
            dominance = 0.9
            evaluate_formulas = false
            ignored_numbers = [0, 1, 7]
            """,
        )
        rules, warnings = load_config(config)
        assert warnings == []
        assert rules.complexity_limit == 3
        assert rules.dominance == 0.9
        assert rules.evaluate_formulas is False
        assert rules.is_ignored_number(7.0)
        assert not rules.is_ignored_number(100.0)

    def test_unknown_keys_are_ignored(self, tmp_path: Path):
        # A file written for a later version must not break an older build.
        config = write(
            tmp_path / "tickmark.toml",
            "[rules]\ncomplexity_limit = 3\nsomething_from_the_future = 12",
        )
        rules, _ = load_config(config)
        assert rules.complexity_limit == 3

    def test_invalid_toml_is_an_error(self, tmp_path: Path):
        config = write(tmp_path / "tickmark.toml", "this is = = not toml")
        with pytest.raises(ConfigError) as excinfo:
            load_config(config)
        assert "not valid TOML" in str(excinfo.value)

    def test_a_setting_of_the_wrong_type_falls_back_without_raising(self, tmp_path: Path):
        config = write(tmp_path / "tickmark.toml", '[rules]\ndominance = "high"')
        rules, warnings = load_config(config)
        assert rules == DEFAULT_RULES
        assert warnings


class TestCustomRules:
    def test_a_rule_becomes_a_finding(self, tmp_path: Path, workbook: Path):
        config = write(
            tmp_path / "tickmark.toml",
            """
            [[rule]]
            name = "no-hardcoded-vat"
            severity = "high"
            pattern = '\\*\\s*1\\.075'
            summary = "Hardcoded VAT rate"
            explanation = "The rate lives in Rates!B1."
            """,
        )
        rules, warnings = load_config(config)
        assert warnings == []

        with open_workbook(workbook) as wb:
            findings = [f for f in run_audit(wb, rules=rules).findings if f.check == "custom-rule"]

        assert findings
        assert findings[0].severity is Severity.HIGH
        # Grouping collapses the run and appends the count, as it does for the
        # built-in checks: a house rule matching six rows is one finding.
        assert findings[0].summary.startswith("Hardcoded VAT rate")
        assert findings[0].span == 6
        assert findings[0].explanation == "The rate lives in Rates!B1."
        assert findings[0].context["rule"] == "no-hardcoded-vat"

    def test_a_rule_can_match_values_rather_than_formulas(self, tmp_path: Path, workbook: Path):
        config = write(
            tmp_path / "tickmark.toml",
            """
            [[rule]]
            name = "no-todo-left"
            target = "value"
            pattern = 'TODO|FIXME'
            summary = "Unfinished note"
            """,
        )
        rules, _ = load_config(config)
        with open_workbook(workbook) as wb:
            findings = [f for f in run_audit(wb, rules=rules).findings if f.check == "custom-rule"]
        assert [f.coordinate for f in findings] == ["D2"]

    def test_no_rules_means_the_check_does_not_run(self, workbook: Path):
        with open_workbook(workbook) as wb:
            findings = run_audit(wb, rules=DEFAULT_RULES).findings
        assert not [f for f in findings if f.check == "custom-rule"]


class TestBadRulesCostOnlyThemselves:
    """Each of these must warn, skip the offending rule, and keep the good one."""

    @pytest.mark.parametrize(
        ("body", "fragment"),
        [
            ('name = "x"\nsummary = "s"', "has no pattern"),
            ('name = "x"\npattern = "p"', "has no summary"),
            ('pattern = "p"\nsummary = "s"', "has no name"),
            (
                'name = "x"\npattern = "p"\nsummary = "s"\nseverity = "catastrophic"',
                "not one of high/medium/low/info",
            ),
            (
                'name = "x"\npattern = "p"\nsummary = "s"\ntarget = "sideways"',
                "not 'formula' or 'value'",
            ),
            (
                'name = "x"\npattern = "(unclosed"\nsummary = "s"',
                "not a valid regular expression",
            ),
        ],
    )
    def test_the_bad_rule_is_named_and_dropped(self, tmp_path: Path, body: str, fragment: str):
        config = write(
            tmp_path / "tickmark.toml",
            f"[[rule]]\n{body}\n\n[[rule]]\nname = \"good\"\npattern = 'TODAY'\nsummary = 'ok'",
        )
        rules, warnings = load_config(config)
        assert any(fragment in w for w in warnings), warnings
        # The valid rule survives: one bad entry costs the user that entry only.
        assert [r.name for r in rules.custom_rules] == ["good"]

    def test_an_uncompilable_pattern_is_rejected_at_load_not_silently_at_run(
        self, tmp_path: Path, workbook: Path
    ):
        """The regression this guards against.

        A pattern that does not compile was dropped at *run* time, where there is
        nowhere to report it — the check runs per cell and cannot warn a thousand
        times. The rule then matched nothing, silently, and the user had no way to
        learn their rule was dead.
        """
        config = write(
            tmp_path / "tickmark.toml",
            '[[rule]]\nname = "broken"\npattern = "(unclosed"\nsummary = "s"',
        )
        rules, warnings = load_config(config)
        assert rules.custom_rules == ()
        assert any("broken" in w for w in warnings)
