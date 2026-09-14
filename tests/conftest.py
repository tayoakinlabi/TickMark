"""Fixture workbooks, built at test time.

Written with openpyxl rather than committed as binaries so the inputs are
readable in the diff — a committed .xlsx is opaque, and a test whose input nobody
can inspect is a test nobody will trust when it fails.

Writing here is fine: these are fixtures being created, not workbooks being
audited. NF 4 governs files Tickmark opens, not files the test suite authors.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName


@pytest.fixture
def simple_workbook(tmp_path: Path) -> Path:
    """A small, correct workbook: a filled-down column plus a total."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws["A1"] = "Item"
    ws["B1"] = "Price"
    ws["C1"] = "WithTax"
    for row in range(2, 6):
        ws[f"A{row}"] = f"item-{row}"
        ws[f"B{row}"] = row * 10
        ws[f"C{row}"] = f"=B{row}*1.075"
    ws["C6"] = "=SUM(C2:C5)"

    path = tmp_path / "simple.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def messy_workbook(tmp_path: Path) -> Path:
    """A workbook carrying the conditions the inventory reports on."""
    wb = Workbook()

    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Header"
    ws["A2"] = 1
    ws["B2"] = "=A2*2"
    ws["B3"] = "=A2+#REF!"
    ws["B4"] = "=NOW()"
    ws["B5"] = "=[1]Budget!A1*2"
    ws.row_dimensions[3].hidden = True
    ws.column_dimensions["D"].hidden = True

    hidden = wb.create_sheet("Hidden")
    hidden["A1"] = "secret"
    hidden.sheet_state = "hidden"

    very = wb.create_sheet("VeryHidden")
    very["A1"] = "more secret"
    very.sheet_state = "veryHidden"

    wb.defined_names.add(DefinedName("TaxRate", attr_text="Data!$A$2"))

    path = tmp_path / "messy.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def not_a_workbook(tmp_path: Path) -> Path:
    path = tmp_path / "notes.txt"
    path.write_text("this is not a workbook")
    return path


@pytest.fixture
def fake_xls(tmp_path: Path) -> Path:
    """A file with the legacy extension. Content is irrelevant: the extension
    alone must be refused, before anything tries to parse it."""
    path = tmp_path / "legacy.xls"
    path.write_bytes(b"\xd0\xcf\x11\xe0garbage")
    return path
